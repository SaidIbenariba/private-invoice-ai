"""Excel workbook a finance team can open directly: summary, invoices, line items, review queue, accuracy."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .validate import Result

HEAD = PatternFill("solid", fgColor="1F2A44")
OK = PatternFill("solid", fgColor="DCF5E3")
FLAG = PatternFill("solid", fgColor="FDE2E1")
MONEY = "#,##0.00"


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int] | None = None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths or [18] * len(headers), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def write_workbook(path: Path, files: list[str], results: list[Result], meta: list[dict], summary: dict | None) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    n = len(results)
    approved = sum(r.status == "APPROVED" for r in results)
    rows = [["Documents processed", n], ["Auto-approved", approved], ["Sent to review", n - approved],
            ["Processing", "100% local (Ollama), no data sent to any cloud API"]]
    if summary:
        rows += [[], ["Measured against ground truth", ""],
                 ["Header field accuracy", f"{summary['header_field_accuracy']:.1%}"],
                 ["Fully correct extractions", f"{summary['fully_correct_extractions']} / {n}"],
                 ["Auto-approve precision", f"{summary['auto_approve_precision']:.1%}"],
                 ["Errors reaching approval", summary["errors_reaching_approval"]],
                 ["Planted anomalies caught", f"{summary['anomalies_caught']} / {summary['anomalies_total']}"],
                 ["Avg seconds / digital PDF", summary["avg_seconds"]["text"]],
                 ["Avg seconds / photo or scan", summary["avg_seconds"]["vision"]]]
    ws = _sheet(wb, "Summary", ["Metric", "Value"], rows, [34, 60])
    ws["B2"].alignment = Alignment(horizontal="left")

    inv_rows = []
    for f, r, m in zip(files, results, meta):
        i = r.invoice
        inv_rows.append([f, r.status, "; ".join(x.message for x in r.issues), i.supplier_name, i.supplier_tax_id,
                         i.invoice_number, i.invoice_date, i.due_date, i.currency, i.subtotal, i.tax_rate,
                         i.tax_amount, i.total, m.get("route"), m.get("seconds")])
    ws = _sheet(wb, "Invoices", ["File", "Status", "Issues", "Supplier", "Tax ID", "Invoice no", "Invoice date",
                                 "Due date", "Currency", "Subtotal", "Tax %", "Tax", "Total", "Path", "Seconds"],
                inv_rows, [14, 11, 48, 28, 18, 14, 12, 12, 9, 12, 7, 11, 12, 8, 8])
    for row in ws.iter_rows(min_row=2):
        row[1].fill = OK if row[1].value == "APPROVED" else FLAG
        for c in (row[9], row[11], row[12]):
            c.number_format = MONEY

    li_rows = [[f, r.invoice.supplier_name, r.invoice.invoice_number, li.description, li.quantity, li.unit_price, li.amount]
               for f, r in zip(files, results) for li in r.invoice.line_items]
    ws = _sheet(wb, "Line items", ["File", "Supplier", "Invoice no", "Description", "Qty", "Unit price", "Amount"],
                li_rows, [14, 28, 14, 36, 8, 12, 12])
    for row in ws.iter_rows(min_row=2):
        row[5].number_format = row[6].number_format = MONEY

    rq = [[f, r.invoice.supplier_name, r.invoice.invoice_number, r.invoice.total, x.code, x.message]
          for f, r in zip(files, results) for x in r.issues]
    _sheet(wb, "Review queue", ["File", "Supplier", "Invoice no", "Total", "Check", "Why it needs a human"],
           rq, [14, 28, 14, 12, 22, 70])

    if summary:
        _sheet(wb, "Accuracy", ["Field", "Accuracy"],
               [[k, f"{v:.1%}"] for k, v in summary["field_accuracy"].items()], [22, 12])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
